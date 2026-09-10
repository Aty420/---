from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADERS = ["任务编号", "源商品模板", "新标题", "SKU名称", "新首图"]
MAP = {
    "任务编号": "task_code",
    "源商品模板": "template_name",
    "新标题": "new_title",
    "SKU名称": "sku_name",
    "新首图": "cover_image",
}


def import_tasks(path: Path) -> list[dict]:
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    headers = {str(c.value).strip(): i for i, c in enumerate(ws[1], 1) if c.value}
    missing = [h for h in HEADERS if h not in headers]
    if missing:
        raise ValueError("缺少必要列：" + "、".join(missing))

    rows = []
    for r in range(2, ws.max_row + 1):
        values = {MAP[h]: ws.cell(r, headers[h]).value for h in HEADERS}
        if not any(values.values()):
            continue
        values = {k: (str(v).strip() if v is not None else "") for k, v in values.items()}
        if not values["task_code"]:
            raise ValueError(f"第 {r} 行缺少任务编号")
        if not values["template_name"]:
            raise ValueError(f"第 {r} 行缺少源商品模板")
        if not values["new_title"]:
            raise ValueError(f"第 {r} 行缺少新标题")
        if not values["sku_name"]:
            raise ValueError(f"第 {r} 行缺少 SKU 名称")
        if not values["cover_image"]:
            raise ValueError(f"第 {r} 行缺少新首图路径")

        image = Path(values["cover_image"])
        if not image.is_absolute():
            image = (path.parent / image).resolve()
        values["cover_image"] = str(image)
        rows.append(values)
    return rows


def create_sample(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "相似品任务"
    ws.append(HEADERS)
    ws.append(["TASK-001", "心相印悬挂抽纸18提", "心相印云感柔肤悬挂抽纸家庭囤货装", "18提家庭装", r"D:\商品素材\001.jpg"])
    ws.append(["TASK-002", "心相印悬挂抽纸18提", "心相印悬挂式抽纸整箱囤货装", "三箱18提", r"D:\商品素材\002.jpg"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="356DFF")
        cell.alignment = Alignment(horizontal="center")
    widths = [18, 24, 52, 24, 46]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
